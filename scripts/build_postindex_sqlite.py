import re
import sqlite3
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = Path("/home/maks_gv/projects/AgroExcelFlow/data 1_5_0/postindex.xlsx")
OUTPUT_PATH = PROJECT_ROOT / "data" / "postindex.sqlite"
SHEET_NAME = None

SOURCE_COLUMNS = {
    "Область": "region_uk",
    "Район (старий)": "district_old_uk",
    "Район (новий)": "district_new_uk",
    "Населений пункт": "settlement_uk",
    "Поштовий індекс (Postal code)": "postal_code",
    "Region (Oblast)": "region_en",
    "District new (Raion new)": "district_new_en",
    "Settlement": "settlement_en",
    "Вiддiлення зв`язку": "post_office_uk",
    "Post office": "post_office_en",
    "Поштовий індекс відділення зв`язку (Post code of post office)": "post_office_postal_code",
}

REQUIRED_COLUMNS = tuple(SOURCE_COLUMNS.keys())


def clean_text(value):
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


def normalize_postal_code(value):
    if value is None:
        return None

    if isinstance(value, int):
        code = str(value)
    elif isinstance(value, float) and value.is_integer():
        code = str(int(value))
    else:
        code = clean_text(value).replace(" ", "")
        if re.fullmatch(r"\d+\.0+", code):
            code = code.split(".", 1)[0]

    digits = re.sub(r"\D", "", code)
    if not digits or len(digits) > 5:
        return None
    return digits.zfill(5)


def get_worksheet(workbook, sheet_name):
    if sheet_name is None:
        return workbook.active
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f'У файлі немає листа "{sheet_name}"')
    return workbook[sheet_name]


def build_column_indexes(header_row):
    headers = [clean_text(value) for value in header_row]
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing_columns:
        missing = ", ".join(missing_columns)
        raise ValueError(f"У файлі відсутні обов'язкові колонки: {missing}")

    return {SOURCE_COLUMNS[header]: headers.index(header) for header in REQUIRED_COLUMNS}


def read_source_rows(source_path, sheet_name=None):
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        worksheet = get_worksheet(workbook, sheet_name)
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise ValueError("Файл довідника порожній")

        column_indexes = build_column_indexes(header_row)
        imported_rows = []
        skipped_rows = []

        for source_row_number, row in enumerate(rows, start=2):
            postal_code = normalize_postal_code(row[column_indexes["postal_code"]])
            if postal_code is None:
                skipped_rows.append(source_row_number)
                continue

            post_office_postal_code = normalize_postal_code(
                row[column_indexes["post_office_postal_code"]]
            )

            imported_rows.append(
                {
                    "source_row": source_row_number,
                    "postal_code": postal_code,
                    "region_uk": clean_text(row[column_indexes["region_uk"]]),
                    "district_old_uk": clean_text(row[column_indexes["district_old_uk"]]),
                    "district_new_uk": clean_text(row[column_indexes["district_new_uk"]]),
                    "settlement_uk": clean_text(row[column_indexes["settlement_uk"]]),
                    "region_en": clean_text(row[column_indexes["region_en"]]),
                    "district_new_en": clean_text(row[column_indexes["district_new_en"]]),
                    "settlement_en": clean_text(row[column_indexes["settlement_en"]]),
                    "post_office_uk": clean_text(row[column_indexes["post_office_uk"]]),
                    "post_office_en": clean_text(row[column_indexes["post_office_en"]]),
                    "post_office_postal_code": post_office_postal_code or "",
                }
            )

        return imported_rows, skipped_rows
    finally:
        workbook.close()


def recreate_schema(connection):
    connection.executescript(
        """
        DROP TABLE IF EXISTS metadata;
        DROP TABLE IF EXISTS postal_code_conflicts;
        DROP TABLE IF EXISTS postal_code_map;
        DROP TABLE IF EXISTS postal_codes;

        CREATE TABLE postal_codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source_row INTEGER NOT NULL,
            postal_code TEXT NOT NULL CHECK(length(postal_code) = 5),
            region_uk TEXT NOT NULL,
            district_old_uk TEXT NOT NULL,
            district_new_uk TEXT NOT NULL,
            settlement_uk TEXT NOT NULL,
            region_en TEXT NOT NULL,
            district_new_en TEXT NOT NULL,
            settlement_en TEXT NOT NULL,
            post_office_uk TEXT NOT NULL,
            post_office_en TEXT NOT NULL,
            post_office_postal_code TEXT NOT NULL
        );

        CREATE INDEX idx_postal_codes_postal_code
            ON postal_codes(postal_code);

        CREATE TABLE postal_code_map (
            postal_code TEXT PRIMARY KEY CHECK(length(postal_code) = 5),
            region_uk TEXT NOT NULL,
            district_old_uk TEXT NOT NULL,
            district_new_uk TEXT NOT NULL,
            status TEXT NOT NULL CHECK(status IN ('ok', 'conflict')),
            source_rows INTEGER NOT NULL,
            settlements_count INTEGER NOT NULL,
            district_variants INTEGER NOT NULL
        );

        CREATE TABLE postal_code_conflicts (
            postal_code TEXT NOT NULL CHECK(length(postal_code) = 5),
            region_uk TEXT NOT NULL,
            district_old_uk TEXT NOT NULL,
            district_new_uk TEXT NOT NULL,
            source_rows INTEGER NOT NULL,
            PRIMARY KEY (
                postal_code,
                region_uk,
                district_old_uk,
                district_new_uk
            )
        );

        CREATE TABLE metadata (
            key TEXT PRIMARY KEY,
            value TEXT NOT NULL
        );
        """
    )


def insert_source_rows(connection, rows):
    connection.executemany(
        """
        INSERT INTO postal_codes (
            source_row,
            postal_code,
            region_uk,
            district_old_uk,
            district_new_uk,
            settlement_uk,
            region_en,
            district_new_en,
            settlement_en,
            post_office_uk,
            post_office_en,
            post_office_postal_code
        )
        VALUES (
            :source_row,
            :postal_code,
            :region_uk,
            :district_old_uk,
            :district_new_uk,
            :settlement_uk,
            :region_en,
            :district_new_en,
            :settlement_en,
            :post_office_uk,
            :post_office_en,
            :post_office_postal_code
        )
        """,
        rows,
    )


def build_postal_code_map(rows):
    rows_by_postal_code = defaultdict(list)
    district_variants_by_postal_code = defaultdict(Counter)
    settlements_by_postal_code = defaultdict(set)

    for row in rows:
        postal_code = row["postal_code"]
        district_key = (
            row["region_uk"],
            row["district_old_uk"],
            row["district_new_uk"],
        )

        rows_by_postal_code[postal_code].append(row)
        district_variants_by_postal_code[postal_code][district_key] += 1

        if row["settlement_uk"]:
            settlements_by_postal_code[postal_code].add(row["settlement_uk"])

    map_rows = []
    conflict_rows = []

    for postal_code in sorted(rows_by_postal_code):
        source_rows = rows_by_postal_code[postal_code]
        source_rows_count = len(source_rows)
        district_variants = district_variants_by_postal_code[postal_code]
        settlements_count = len(settlements_by_postal_code[postal_code])
        region_uk = get_single_value(row["region_uk"] for row in source_rows)
        district_old_uk = get_single_value(row["district_old_uk"] for row in source_rows)
        district_new_uk = get_single_value(row["district_new_uk"] for row in source_rows)

        if len(district_variants) == 1:
            status = "ok"
        else:
            status = "conflict"

            for district_key, rows_count in sorted(district_variants.items()):
                conflict_rows.append(
                    {
                        "postal_code": postal_code,
                        "region_uk": district_key[0],
                        "district_old_uk": district_key[1],
                        "district_new_uk": district_key[2],
                        "source_rows": rows_count,
                    }
                )

        map_rows.append(
            {
                "postal_code": postal_code,
                "region_uk": region_uk,
                "district_old_uk": district_old_uk,
                "district_new_uk": district_new_uk,
                "status": status,
                "source_rows": source_rows_count,
                "settlements_count": settlements_count,
                "district_variants": len(district_variants),
            }
        )

    return map_rows, conflict_rows


def get_single_value(values):
    unique_values = sorted(set(values))
    if len(unique_values) == 1:
        return unique_values[0]
    return ""


def insert_postal_code_map(connection, map_rows, conflict_rows):
    connection.executemany(
        """
        INSERT INTO postal_code_map (
            postal_code,
            region_uk,
            district_old_uk,
            district_new_uk,
            status,
            source_rows,
            settlements_count,
            district_variants
        )
        VALUES (
            :postal_code,
            :region_uk,
            :district_old_uk,
            :district_new_uk,
            :status,
            :source_rows,
            :settlements_count,
            :district_variants
        )
        """,
        map_rows,
    )

    connection.executemany(
        """
        INSERT INTO postal_code_conflicts (
            postal_code,
            region_uk,
            district_old_uk,
            district_new_uk,
            source_rows
        )
        VALUES (
            :postal_code,
            :region_uk,
            :district_old_uk,
            :district_new_uk,
            :source_rows
        )
        """,
        conflict_rows,
    )


def insert_metadata(connection, source_path, imported_rows, skipped_rows, map_rows, conflict_rows):
    metadata = {
        "source_path": str(source_path),
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "imported_source_rows": str(len(imported_rows)),
        "skipped_source_rows": str(len(skipped_rows)),
        "unique_postal_codes": str(len(map_rows)),
        "conflict_postal_codes": str(
            sum(1 for row in map_rows if row["status"] == "conflict")
        ),
        "conflict_variants": str(len(conflict_rows)),
    }

    connection.executemany(
        "INSERT INTO metadata (key, value) VALUES (?, ?)",
        sorted(metadata.items()),
    )


def write_sqlite_database(output_path, source_path, rows, skipped_rows):
    output_path.parent.mkdir(parents=True, exist_ok=True)

    map_rows, conflict_rows = build_postal_code_map(rows)

    with sqlite3.connect(output_path) as connection:
        recreate_schema(connection)
        insert_source_rows(connection, rows)
        insert_postal_code_map(connection, map_rows, conflict_rows)
        insert_metadata(connection, source_path, rows, skipped_rows, map_rows, conflict_rows)

    return map_rows, conflict_rows


def main():
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(f"Файл довідника не знайдено: {SOURCE_PATH}")

    rows, skipped_rows = read_source_rows(SOURCE_PATH, SHEET_NAME)
    map_rows, conflict_rows = write_sqlite_database(
        OUTPUT_PATH,
        SOURCE_PATH,
        rows,
        skipped_rows,
    )

    conflict_postal_codes = sum(1 for row in map_rows if row["status"] == "conflict")

    print(f"SQLite файл створено: {OUTPUT_PATH}")
    print(f"Імпортовано рядків: {len(rows)}")
    print(f"Пропущено рядків без валідного індексу: {len(skipped_rows)}")
    print(f"Унікальних поштових індексів: {len(map_rows)}")
    print(f"Індексів з конфліктними районами: {conflict_postal_codes}")

    if skipped_rows:
        preview = ", ".join(str(row_number) for row_number in skipped_rows[:10])
        print(f"Перші пропущені рядки: {preview}")

    if conflict_postal_codes:
        print("Деталі конфліктів записано в таблицю postal_code_conflicts")


if __name__ == "__main__":
    main()
